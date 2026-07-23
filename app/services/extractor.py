import asyncio
import re
from pathlib import Path

from docx import Document
from pypdf import PdfReader


class ExtractionError(ValueError):
    pass


SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt"}


def normalize_text(text: str) -> str:
    text = text.casefold()
    for apostrophe in ("ʻ", "ʼ", "‘", "’", "`", "´"):
        text = text.replace(apostrophe, "'")
    text = re.sub(r"[^\w\s']", " ", text, flags=re.UNICODE)
    return re.sub(r"\s+", " ", text).strip()


def _extract_pdf(path: Path) -> str:
    try:
        reader = PdfReader(str(path))
        return "\n".join(page.extract_text() or "" for page in reader.pages)
    except Exception as exc:
        raise ExtractionError("PDF faylni o‘qib bo‘lmadi.") from exc


def _extract_docx(path: Path) -> str:
    try:
        document = Document(str(path))
        paragraphs = [paragraph.text for paragraph in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                paragraphs.append(" | ".join(cell.text for cell in row.cells))
        return "\n".join(paragraphs)
    except Exception as exc:
        raise ExtractionError("DOCX faylni o‘qib bo‘lmadi.") from exc


def _extract_txt(path: Path) -> str:
    raw = path.read_bytes()
    for encoding in ("utf-8-sig", "utf-8", "cp1251", "latin-1"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ExtractionError("TXT fayl kodirovkasini aniqlab bo‘lmadi.")


async def extract_text(path: Path, extension: str) -> str:
    extension = extension.casefold()
    if extension not in SUPPORTED_EXTENSIONS:
        raise ExtractionError("Faqat PDF, DOCX va TXT fayllar qabul qilinadi.")

    if extension == ".pdf":
        text = await asyncio.to_thread(_extract_pdf, path)
    elif extension == ".docx":
        text = await asyncio.to_thread(_extract_docx, path)
    else:
        text = await asyncio.to_thread(_extract_txt, path)

    text = text.replace("\x00", "").strip()
    if len(normalize_text(text).split()) < 20:
        raise ExtractionError(
            "Fayldan yetarli matn topilmadi. Skan qilingan PDF hozircha qo‘llanmaydi."
        )
    return text
